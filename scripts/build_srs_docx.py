#!/usr/bin/env python3
"""Build SRS_Electricity_Load_Forecasting_IEEE.docx from the markdown project report with Word styles."""

from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

IMG_MD = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")


def set_document_default_font(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)


def add_runs_with_bold(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = "Times New Roman"
        else:
            run = paragraph.add_run(part)
            run.font.name = "Times New Roman"


def add_paragraph(doc: Document, text: str, *, style: str | None = None) -> None:
    if not text.strip():
        return
    p = doc.add_paragraph(style=style)
    add_runs_with_bold(p, text.strip())


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            val = row[j] if j < len(row) else ""
            p = cell.paragraphs[0]
            add_runs_with_bold(p, val)
    doc.add_paragraph()


def parse_md_table(block: list[str]) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = 0
    while i < len(block):
        line = block[i].strip()
        if not line.startswith("|"):
            break
        if re.match(r"^\|\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*$", line):
            i += 1
            continue
        raw_cells = line.split("|")
        cells = [c.strip() for c in raw_cells]
        if cells and cells[0] == "":
            cells = cells[1:]
        if cells and cells[-1] == "":
            cells = cells[:-1]
        if cells:
            rows.append(cells)
        i += 1
    return rows, i


def add_markdown_image(doc: Document, md_dir: Path, _alt: str, rel_path: str) -> None:
    pic = (md_dir / rel_path.strip()).resolve()
    if not pic.is_file():
        add_paragraph(doc, f"[Missing image file: {rel_path}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(pic), width=Inches(6.45))
    doc.add_paragraph()
    # Alt text is not duplicated here; the SRS markdown supplies *Figure …* captions after the image.


def md_to_docx(md_path: Path, out_path: Path) -> None:
    md_dir = md_path.parent
    raw = md_path.read_text(encoding="utf-8")
    raw = re.sub(r"<!--.*?-->\s*", "", raw, flags=re.DOTALL)
    lines = raw.split("\n")

    # Start body at Abstract (skip duplicate title block in .md)
    start = 0
    for i, line in enumerate(lines):
        if line.strip() == "### Abstract":
            start = i
            break
    lines = lines[start:]

    doc = Document()
    set_document_default_font(doc)

    # Title page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("National University of Computer & Emerging Sciences, Karachi")
    r.bold = True
    r.font.size = Pt(12)
    r.font.name = "Times New Roman"

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Project Report")
    r.bold = True
    r.font.size = Pt(16)
    r.font.name = "Times New Roman"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Deep Learning Electricity Load Forecasting System")
    r.bold = True
    r.font.size = Pt(14)
    r.font.name = "Times New Roman"

    doc.add_paragraph()
    for line in (
        "Course: Software Engineering",
        "Semester: Spring 2026",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Prepared By")
    r.bold = True
    r.font.name = "Times New Roman"
    for name in ("Abdul Wasay — 23L-0658", "Haseeb Asif — 23K-0539"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(name)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    doc.add_page_break()

    idx = 0
    while idx < len(lines):
        line = lines[idx]

        if line.strip() == "---":
            idx += 1
            continue

        if line.strip().startswith("|") and line.count("|") >= 2:
            rows, consumed = parse_md_table(lines[idx:])
            if rows:
                add_table(doc, rows)
            idx += consumed
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            idx += 1
            continue

        if line.startswith("### "):
            rest = line[4:].strip()
            if rest in ("*End of SRS*", "End of SRS", "*End of Project Report*", "End of Project Report"):
                para = doc.add_paragraph()
                run = para.add_run("End of Project Report")
                run.italic = True
                run.font.name = "Times New Roman"
            else:
                doc.add_heading(rest, level=2)
            idx += 1
            continue

        if line.startswith("# "):  # should not appear after slicing
            doc.add_heading(line[2:].strip(), level=1)
            idx += 1
            continue

        if line.startswith("```"):
            idx += 1
            buf: list[str] = []
            while idx < len(lines) and not lines[idx].strip().startswith("```"):
                buf.append(lines[idx])
                idx += 1
            idx += 1
            para = doc.add_paragraph()
            run = para.add_run("\n".join(buf))
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            continue

        m_img = IMG_MD.match(line.strip())
        if m_img:
            add_markdown_image(doc, md_dir, m_img.group(1), m_img.group(2))
            idx += 1
            continue

        stripped = line.lstrip()

        if "\\(" in line or "\\mathrm" in line:
            add_paragraph(
                doc,
                "Skill score (vs seasonal naive): 1 − (MAE_LSTM / MAE_naive) on the same "
                "test windows (higher is better when LSTM beats naive).",
            )
            idx += 1
            continue

        if stripped.startswith("*") and stripped.endswith("*") and len(stripped) > 2:
            p = doc.add_paragraph()
            r = p.add_run(stripped.strip("*"))
            r.italic = True
            idx += 1
            continue

        if stripped.startswith("- "):
            add_paragraph(doc, stripped[2:].strip(), style="List Bullet")
            idx += 1
            continue

        if stripped.startswith("✅"):
            add_paragraph(doc, stripped[1:].strip(), style="List Bullet")
            idx += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            add_paragraph(doc, stripped.strip(), style="List Number")
            idx += 1
            continue

        if line.strip():
            add_paragraph(doc, line.strip())
        idx += 1

    doc.save(out_path)


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    md = root / "SRS_Electricity_Load_Forecasting_IEEE.md"
    out = root / "SRS_Electricity_Load_Forecasting_IEEE.docx"
    if not md.is_file():
        raise SystemExit(f"Missing {md}")
    gen = root / "scripts" / "generate_architecture_diagram.py"
    if gen.is_file():
        subprocess.run([sys.executable, str(gen)], check=True, cwd=str(root))
    md_to_docx(md, out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
